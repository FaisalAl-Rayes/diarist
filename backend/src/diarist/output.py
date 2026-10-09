from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from .models import Transcript, Utterance, Word


class OutputError(RuntimeError):
    pass


def format_timestamp(seconds: float) -> str:
    total = max(0, int(seconds))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def write_json(transcript: Transcript, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(transcript.to_dict(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def read_json(path: Path) -> Transcript:
    """Load a transcript JSON written by this application."""
    path = path.expanduser().resolve()
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        utterances = []
        for item in data["utterances"]:
            words = [
                Word(
                    start=float(word["start"]),
                    end=float(word["end"]),
                    text=str(word["text"]),
                    speaker=word.get("speaker"),
                    probability=(
                        float(word["probability"])
                        if word.get("probability") is not None
                        else None
                    ),
                )
                for word in item.get("words", [])
            ]
            utterances.append(
                Utterance(
                    start=float(item["start"]),
                    end=float(item["end"]),
                    speaker=str(item["speaker"]),
                    text=str(item["text"]),
                    words=words,
                    id=int(item["id"]) if item.get("id") is not None else None,
                )
            )
        return Transcript(
            source=Path(data["source"]),
            language=str(data["language"]),
            requested_language=data.get("requested_language"),
            duration=float(data["duration"]),
            model=str(data["models"]["transcription"]),
            diarization_model=str(data["models"]["diarization"]),
            requested_speakers=(
                int(data["requested_speakers"])
                if data.get("requested_speakers") is not None
                else None
            ),
            speakers=[str(speaker) for speaker in data["speakers"]],
            utterances=utterances,
        )
    except FileNotFoundError as exc:
        raise OutputError(f"Transcript JSON not found: {path}") from exc
    except (json.JSONDecodeError, KeyError, TypeError, ValueError) as exc:
        raise OutputError(f"Invalid transcript JSON: {path}") from exc


def _set_cell_margins(cell: object, *, top: int, start: int, bottom: int, end: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table: object, widths_dxa: tuple[int, ...]) -> None:
    """Set matching fixed widths on the table, grid, and every cell."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    table_xml = table._tbl
    properties = table_xml.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(sum(widths_dxa)))
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    for grid_column, width in zip(table_xml.tblGrid.gridCol_lst, widths_dxa, strict=True):
        grid_column.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths_dxa, strict=True):
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(width))


def _add_page_number(paragraph: object) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])


def write_docx(transcript: Transcript, path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    path.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15
    for name, size, before, after in (
        ("Title", 25, 0, 6),
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
    ):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(31, 77, 120)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.text = "TRANSCRIPTION"
    header.style = styles["Caption"]
    header.runs[0].font.color.rgb = RGBColor(96, 105, 115)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    label = footer.add_run("Page ")
    label.font.size = Pt(9)
    label.font.color.rgb = RGBColor(96, 105, 115)
    _add_page_number(footer)

    title = document.add_paragraph(style="Title")
    title.add_run(transcript.source.stem.replace("_", " ").replace("-", " ").title())
    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Multi-speaker transcription")
    run.italic = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(96, 105, 115)

    metadata = document.add_table(rows=2, cols=3)
    metadata.autofit = False
    widths_dxa = (3120, 3120, 3120)
    _set_table_geometry(metadata, widths_dxa)
    values = (
        ("SOURCE", transcript.source.name),
        ("DURATION", format_timestamp(transcript.duration)),
        ("SPEAKERS", str(len(transcript.speakers))),
        ("LANGUAGE", transcript.language.upper()),
        ("MODEL", "Whisper large-v3"),
        ("CREATED", datetime.now().astimezone().strftime("%Y-%m-%d %H:%M")),
    )
    for index, cell in enumerate(cell for row in metadata.rows for cell in row.cells):
        _set_cell_margins(cell, top=90, start=120, bottom=90, end=120)
        label_text, value = values[index]
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        label_run = paragraph.add_run(label_text + "\n")
        label_run.bold = True
        label_run.font.size = Pt(8)
        label_run.font.color.rgb = RGBColor(78, 121, 167)
        value_run = paragraph.add_run(value)
        value_run.font.size = Pt(10)

    document.add_heading("Transcription", level=1)
    colors = [
        RGBColor(31, 77, 120),
        RGBColor(147, 75, 46),
        RGBColor(72, 116, 82),
        RGBColor(105, 78, 132),
        RGBColor(142, 109, 31),
        RGBColor(50, 112, 125),
    ]
    speaker_colors = {
        speaker: colors[index % len(colors)] for index, speaker in enumerate(transcript.speakers)
    }
    for utterance in transcript.utterances:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(7)
        marker = paragraph.add_run(f"[{format_timestamp(utterance.start)}]  {utterance.speaker}  ")
        marker.bold = True
        marker.font.size = Pt(9)
        marker.font.color.rgb = speaker_colors.get(utterance.speaker, RGBColor(60, 60, 60))
        text_run = paragraph.add_run(utterance.text)
        text_run.font.size = Pt(10.5)

    document.save(str(path))
